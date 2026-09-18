import openpyxl
import mysql.connector

# Connect to MySQL
db = mysql.connector.connect(
    host="localhost",
    user="root",
    password="password",
    database="smart_canteen"
)

cursor = db.cursor()

# Open Excel file
workbook = openpyxl.load_workbook("csit details(1).xlsx")
sheet = workbook.active

# Data starts from row 3
for row in range(3, sheet.max_row + 1):

    registration_no = sheet.cell(row=row, column=2).value
    student_name = sheet.cell(row=row, column=3).value

    if registration_no is None or student_name is None:
        continue

    registration_no = str(registration_no).strip()
    student_name = str(student_name).strip()

    cursor.execute(
        """
        UPDATE users
        SET student_name = %s
        WHERE college_id = %s
        AND user_type = 'student'
        """,
        (student_name, registration_no)
    )

    print("Updated:", registration_no, "-", student_name)

db.commit()
cursor.close()
db.close()

print("Student names updated successfully!")